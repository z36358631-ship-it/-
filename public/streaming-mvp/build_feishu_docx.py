from html.parser import HTMLParser
from pathlib import Path
from html import unescape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import zipfile, re, hashlib, json

ROOT = Path(__file__).resolve().parent
HTML = ROOT / 'prd.html'
OUT = ROOT / '盖世远程PRD-V4.11-Word图片内嵌版.docx'

class Node:
    def __init__(self, tag=None, attrs=None, parent=None, text=None):
        self.tag, self.attrs, self.parent, self.text = tag, dict(attrs or []), parent, text
        self.children = []
    def add(self, n): self.children.append(n); return n

class Parser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node('root'); self.cur = self.root
    def handle_starttag(self, tag, attrs):
        n = self.cur.add(Node(tag, attrs, self.cur));
        if tag not in {'img','br','hr','meta','link','input'}: self.cur=n
    def handle_startendtag(self, tag, attrs): self.cur.add(Node(tag, attrs, self.cur))
    def handle_endtag(self, tag):
        c=self.cur
        while c is not self.root and c.tag != tag: c=c.parent
        if c is not self.root and c.parent: self.cur=c.parent
    def handle_data(self, data):
        if data: self.cur.add(Node(text=data, parent=self.cur))

html = HTML.read_text(encoding='utf-8')
p = Parser(); p.feed(html)

def first(n, tag):
    if n.tag == tag: return n
    for c in n.children:
        if c.text is None:
            v=first(c,tag)
            if v: return v
    return None
main = first(p.root,'main')
assert main is not None

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.55); sec.bottom_margin = Inches(0.55); sec.left_margin = Inches(0.65); sec.right_margin = Inches(0.65)
styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size = Pt(9.5)
for name, size in [('Title',20),('Heading 1',16),('Heading 2',13),('Heading 3',11),('Heading 4',10.5),('Heading 5',10)]:
    st=styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def text_of(n):
    if n.text is not None: return n.text
    return ''.join(text_of(c) for c in n.children)

def img_path(node):
    src=node.attrs.get('src','')
    if src.startswith('data:'): return None
    return (HTML.parent / src).resolve()

def add_image(run, node, in_cell=False):
    path=img_path(node)
    if not path or not path.exists():
        run.add_text('[图片缺失]')
        return
    try:
        with Image.open(path) as im: w,h=im.size
        max_w = 3.15 if in_cell else 6.2
        max_h = 4.0 if in_cell else 5.2
        scale=min(max_w/(w/96), max_h/(h/96), 1.0)
        run.add_picture(str(path), width=Inches((w/96)*scale), height=Inches((h/96)*scale))
    except Exception:
        run.add_picture(str(path), width=Inches(3.15 if in_cell else 5.8))

def inline(paragraph, nodes, bold=False, italic=False, in_cell=False):
    for n in nodes:
        if n.text is not None:
            t=n.text
            if t: 
                r=paragraph.add_run(t); r.bold=bold; r.italic=italic
            continue
        tag=n.tag
        if tag=='br': paragraph.add_run().add_break(); continue
        if tag=='img': add_image(paragraph.add_run(),n,in_cell); continue
        if tag in ('strong','b'): inline(paragraph,n.children,True,italic,in_cell); continue
        if tag in ('em','i'): inline(paragraph,n.children,bold,True,in_cell); continue
        if tag=='code':
            r=paragraph.add_run(text_of(n)); r.font.name='Consolas'; r.font.size=Pt(8.5); continue
        if tag=='a': inline(paragraph,n.children,bold,italic,in_cell); continue
        inline(paragraph,n.children,bold,italic,in_cell)

def render_cell(cell, node):
    cell.text=''; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; set_cell_margins(cell)
    # Render every child in source order; filtering tags discarded strong/em/a.
    para=cell.paragraphs[0]
    para.paragraph_format.space_after=Pt(2)
    inline(para,node.children,in_cell=True)

def render_table(node):
    trs=[]
    def walk(n):
        if n.tag=='tr': trs.append(n); return
        for c in n.children:
            if c.text is None: walk(c)
    walk(node)
    if not trs: return
    maxcols=max(sum(1 for c in tr.children if c.text is None and c.tag in ('td','th')) for tr in trs)
    table=doc.add_table(rows=0, cols=maxcols); table.style='Table Grid'; table.autofit=True
    for ri,tr in enumerate(trs):
        cells=[c for c in tr.children if c.text is None and c.tag in ('td','th')]
        row=table.add_row().cells
        for ci in range(maxcols):
            c=row[ci]
            if ci < len(cells):
                render_cell(c,cells[ci])
                if ri==0 or cells[ci].tag=='th': set_cell_shading(c,'E9EDF2')
            else: c.text=''
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_block(node):
    tag=node.tag
    if tag in ('h1','h2','h3','h4','h5','h6'):
        para=doc.add_paragraph(style={'h1':'Title','h2':'Heading 1','h3':'Heading 2','h4':'Heading 3','h5':'Heading 4','h6':'Heading 5'}[tag]); inline(para,node.children)
    elif tag=='table': render_table(node)
    elif tag in ('p','blockquote'):
        para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(5); inline(para,node.children)
    elif tag in ('ul','ol'):
        for li in node.children:
            if li.text is None and li.tag=='li':
                para=doc.add_paragraph(style='List Bullet' if tag=='ul' else 'List Number'); inline(para,li.children)
    elif tag=='hr': doc.add_paragraph()
    elif tag in ('div','section'):
        for c in node.children:
            if c.text is None: add_block(c)
    else:
        if node.text is None:
            for c in node.children:
                if c.text is None: add_block(c)

# Intro for the import target.
title=doc.add_paragraph(style='Title'); title.add_run('盖世远程独立客户端与移动 App PRD V4.11')
intro=doc.add_paragraph(); intro.add_run('飞书导入版：').bold=True; intro.add_run('本文件已将 PRD 中的本地截图直接内嵌到 Word 文档。请将本 DOCX 文件直接导入飞书，图片不依赖本地 assets 文件夹。')

for child in main.children:
    if child.text is not None: continue
    if child.tag in ('h1','h2','h3','h4','h5','h6','p','table','blockquote','ul','ol','hr'):
        # skip original HTML title h1; the import title above is cleaner.
        if child.tag=='h1': continue
        add_block(child)

# Prevent rows from splitting where possible and set document core props.
for table in doc.tables:
    for row in table.rows:
        trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs: run.font.name='Microsoft YaHei'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.core_properties.title='盖世远程独立客户端与移动 App PRD V4.11'
doc.core_properties.subject='飞书导入版，图片已内嵌'
doc.core_properties.author='盖世远程产品团队'
doc.save(OUT)

# Verify the saved file against every source heading and table cell.
def descendants(node, tag):
    found=[]
    for child in node.children:
        if child.text is None:
            if child.tag == tag: found.append(child)
            found.extend(descendants(child,tag))
    return found

def compact(value): return re.sub(r'\s+', '', value)

saved=Document(OUT)
source_headings=[text_of(n) for n in main.children if n.tag in ('h2','h3','h4','h5','h6')]
saved_headings=[para.text for para in saved.paragraphs if para.style.name.startswith('Heading ')]
assert source_headings == saved_headings, 'Heading content/order changed during export'
source_tables=descendants(main,'table')
assert len(source_tables)==len(saved.tables), 'Table count changed during export'
checked_cells=0
label_counts={'展示说明：':0,'交互说明：':0}
for source_table,saved_table in zip(source_tables,saved.tables):
    source_rows=descendants(source_table,'tr')
    assert len(source_rows)==len(saved_table.rows), 'Table row count changed during export'
    for source_row,saved_row in zip(source_rows,saved_table.rows):
        source_cells=[n for n in source_row.children if n.tag in ('td','th')]
        assert len(source_cells)==len(saved_row.cells), 'Table column count changed during export'
        for source_cell,saved_cell in zip(source_cells,saved_row.cells):
            assert compact(text_of(source_cell)) == compact(saved_cell.text), 'Cell text lost during export: '+text_of(source_cell)[:80]
            expected_bold=''.join(text_of(n) for n in descendants(source_cell,'strong'))
            actual_bold=''.join(r.text for para in saved_cell.paragraphs for r in para.runs if r.bold)
            assert compact(expected_bold)==compact(actual_bold), 'Bold labels lost during export'
            for label in label_counts:
                expected=text_of(source_cell).count(label)
                assert saved_cell.text.count(label)==expected, 'Section label missing: '+label
                label_counts[label]+=expected
            checked_cells+=1
source_image_count=len(descendants(main,'img'))

# Verification report.
with zipfile.ZipFile(OUT) as z:
    media=[n for n in z.namelist() if n.startswith('word/media/')]
    document_xml=z.read('word/document.xml')
assert len(saved.inline_shapes)==source_image_count, 'Image references lost during export'
report={'output':str(OUT),'status':'content-check-passed','headingCount':len(saved_headings),'pageHeadingCount':sum(n.tag=='h4' and text_of(n).startswith(('3.1.','3.2.')) for n in main.children),'sourceHeadingOrderPreserved':True,'tableCount':len(saved.tables),'cellsChecked':checked_cells,'allCellTextPreserved':True,'boldLabelsPreserved':True,'sectionLabels':label_counts,'imageReferences':len(saved.inline_shapes),'embeddedMediaCount':len(media),'mediaFiles':media,'documentHasDrawing':b'<w:drawing' in document_xml,'size':OUT.stat().st_size,'feishuImport':'not-tested-this-revision'}
(ROOT/'feishu-docx-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
