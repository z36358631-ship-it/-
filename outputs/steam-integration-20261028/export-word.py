"""Export the approved PRD with editable tables and embedded figures."""
from pathlib import Path
import re, json, zipfile, urllib.parse, argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT = Path(__file__).resolve().parent
parser = argparse.ArgumentParser()
parser.add_argument('--version', default='V1.8')
VERSION = parser.parse_args().version
OUT = ROOT / ('盖世游戏Steam接入方案' + VERSION + '.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Heading 4','Heading 5']:
    st = doc.styles[name]
    st.font.name = 'Microsoft YaHei'
    st.font.color.rgb = RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.size = Pt(11 if name == 'Normal' else {'Title':22,'Heading 1':17,'Heading 2':15,'Heading 3':13,'Heading 4':12,'Heading 5':11}[name])
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.18
doc.core_properties.title = '盖世游戏 Steam 接入方案'
doc.core_properties.subject = VERSION + ' 产品需求与页面规则'
doc.core_properties.author = '盖世游戏'
doc.core_properties.comments = ''
figures = []
TOKEN = re.compile(r'(!?\[[^\]]*\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)')

def hyperlink(p, label, url):
    if not re.match(r'https?://', url):
        url = 'https://github.com/z36358631-ship-it/-/blob/codex/steam-integration-20261028/outputs/steam-integration-20261028/' + urllib.parse.quote(url)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '175CD3'); pr.append(color)
    r.append(pr); t = OxmlElement('w:t'); t.text = label; r.append(t); link.append(r); p._p.append(link)

def inline(p, value, width=7.2):
    for token in TOKEN.split(value):
        if not token: continue
        if token.startswith('!['):
            alt, url = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', token).groups()
            img = ROOT / 'prd-assets' / urllib.parse.unquote(url.rsplit('/',1)[-1])
            assert img.exists(), img
            w,h = Image.open(img).size
            draw_w = min(width, 3.3 if h>w else width, 4.4*w/h)
            if p.text: p.add_run().add_break()
            shape = p.add_run().add_picture(str(img), width=Inches(draw_w))
            shape._inline.docPr.set('descr', alt)
            figures.append(img.name)
        elif token.startswith('['):
            label,url = re.match(r'\[([^\]]*)\]\(([^)]+)\)',token).groups(); hyperlink(p,label,url)
        elif token.startswith('**'):
            p.add_run(token[2:-2]).bold=True
        elif token.startswith('`'):
            p.add_run(token[1:-1])
        else: p.add_run(token)

def cell_content(cell, value, width):
    for i,part in enumerate(re.split(r'<br\s*/?>', value)):
        p = cell.paragraphs[0] if i==0 else cell.add_paragraph()
        p.paragraph_format.space_after=Pt(4)
        inline(p,part,width-.18)

def table(lines):
    rows = [[s.strip() for s in l.strip().strip('|').split('|')] for l in lines]
    rows = [r for r in rows if not all(re.fullmatch(r'[:\- ]+',x) for x in r)]
    n=len(rows[0]); heads=rows[0]
    if n==2: widths=[1.12,6.08]
    elif heads[0]=='类型': widths=[.7,2.15,1.1,3.25]
    elif heads[0]=='事件': widths=[1.7,1.0,2.4,2.1]
    elif heads[0]=='参数': widths=[1.55,1.1,2.1,2.45]
    elif n==4: widths=[1.1,2.25,1.65,2.2]
    elif n==3: widths=[1.45,2.7,3.05]
    else: widths=[7.2/n]*n
    t=doc.add_table(rows=0, cols=n); t.autofit=False
    t.style='Table Grid'
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for row_index, values in enumerate(rows):
        cells=t.add_row().cells
        for cell,w,value in zip(cells,widths,values):
            cell.width=Inches(w)
            cell_content(cell,value,w)
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if row_index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF2');tcpr.append(shade)
                for p in cell.paragraphs:
                    for run in p.runs: run.bold=True
        if row_index==0:
            repeat=OxmlElement('w:tblHeader');t.rows[-1]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def markdown(text, appendix=False):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1;continue
        if line.startswith('|'):
            group=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                group.append(lines[i]);i+=1
            table(group);continue
        h=re.match(r'^(#{1,6})\s+(.*)',line)
        if h:
            level=len(h[1]); label=h[2]
            if level==1 and not appendix:
                doc.add_paragraph('盖世游戏 Steam 接入方案',style='Title')
                doc.add_paragraph(VERSION + '  ·  产品需求文档')
            elif level==1: doc.add_heading('6.3 第三方 Steam 服务调用免责声明',level=2)
            else: doc.add_heading(label,level=min(level-1 if not appendix else level,5))
        elif line=='---': pass
        else:
            p=doc.add_paragraph();inline(p,line)
        i+=1

source=(ROOT/'Steam接入PRD-Draft.md').read_text(encoding='utf-8-sig')
markdown(source)
doc.add_page_break()
markdown((ROOT/'免责声明.md').read_text(encoding='utf-8-sig'),appendix=True)
actual = re.sub(r'\s+', '', ''.join(n.text or '' for n in doc._element.xpath('.//w:t')))
checked = 0
for content in [source, (ROOT/'免责声明.md').read_text(encoding='utf-8-sig')]:
    for line in content.splitlines():
        line = line.strip()
        if not line or line.startswith('# ') or line == '---': continue
        parts = line.strip('|').split('|') if line.startswith('|') else [re.sub(r'^#{1,6}\s+', '', line)]
        if all(re.fullmatch(r'[\s:\-]+', p) for p in parts): continue
        for part in parts:
            part = re.sub(r'!\[[^\]]*\]\([^)]+\)', '', part)
            part = re.sub(r'\[([^\]]*)\]\([^)]+\)', r'\1', part)
            part = re.sub(r'<br\s*/?>', '', part).replace('**', '').replace('`', '')
            part = re.sub(r'\s+', '', part)
            if part:
                assert part in actual, 'Missing document text: ' + part[:90]
                checked += 1
doc.save(OUT)
with zipfile.ZipFile(OUT) as z:
    media=[n for n in z.namelist() if n.startswith('word/media/')]
    rel=z.read('word/_rels/document.xml.rels').decode()
    assert 'relationships/image' in rel
    assert not re.search(r'Type="[^"]*/image"[^>]*TargetMode="External"',rel)
check={'file':OUT.name,'source':'Steam接入PRD-Draft.md','version':VERSION,'tables':len(doc.tables),'imageOccurrences':len(figures),'embeddedImages':len(media),'allImagesEmbedded':True,'disclaimerIncluded':True,'figures':figures,'renderStatus':'pending','feishuImport':'not-tested'}
check.update({'textSegmentsChecked':checked, 'missingTextSegments':[]})
(ROOT/('Word导出验证-'+VERSION+'.json')).write_text(json.dumps(check,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(check,ensure_ascii=False))
